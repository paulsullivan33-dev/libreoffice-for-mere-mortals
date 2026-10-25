#!/usr/bin/env python3
"""Create the guide's formula fixture: seed-orders.xlsx.

openpyxl writes the formulas but no cached results (it
never calculates). That gap is the point of the fixture:
what does a converter do with a formula whose stored
answer is missing?
"""
from openpyxl import Workbook

wb = Workbook()
ws = wb.active
ws.title = "Orders"
ws.append(["Item", "Packets", "Price", "Line total"])
rows = [
    ("Tomato - Gardener's Delight", 6, 2.50),
    ("Sweet Pea - Spencer Mixed", 4, 1.95),
    ("Courgette - Defender F1", 2, 3.25),
]
for item, packets, price in rows:
    ws.append([item, packets, price])
for r in range(2, 5):
    ws.cell(row=r, column=4, value=f"=B{r}*C{r}")
ws.cell(row=5, column=3, value="Total")
ws.cell(row=5, column=4, value="=SUM(D2:D4)")

wb.save("seed-orders.xlsx")
print("wrote seed-orders.xlsx (formulas, no cached values)")
